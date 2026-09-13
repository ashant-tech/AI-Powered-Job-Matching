

def parse_docx(file_path: str) -> str | None:
    """
    Parse text from a DOCX file.
    """
    try:
        from docx import Document
        doc = Document(file_path)
        text = ""
        
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        return text
    except ImportError:
        print("python-docx not installed. Install with: pip install python-docx")
        return ""
    except Exception as e:
        print(f"Error parsing DOCX: {e}")
        return ""

def parse_docx_bytes(docx_bytes: bytes) -> str:
    """
    Parse text from DOCX bytes.
    """
    try:
        from docx import Document
        import io
        doc_file = io.BytesIO(docx_bytes)
        doc = Document(doc_file)
        text = ""
        
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        return text
    except ImportError:
        print("python-docx not installed. Install with: pip install python-docx")
        return ""
    except Exception as e:
        print(f"Error parsing DOCX bytes: {e}")
        return ""
